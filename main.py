import os
import re
import zipfile
import shutil
from threading import Thread
from datetime import datetime

from kivy.app import App
from kivy.clock import Clock
from kivy.graphics import Color, Line
from kivy.metrics import dp
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.progressbar import ProgressBar
from kivy.uix.scrollview import ScrollView
from kivy.uix.widget import Widget

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage


# ============================================================
# PIL - IMAGE PREVIEWS
# ============================================================

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


# ============================================================
# ANDROID / PYJNIUS
# ============================================================

try:
    from jnius import autoclass
    ANDROID_AVAILABLE = True
except Exception:
    ANDROID_AVAILABLE = False


# ============================================================
# ANDROID STORAGE ACCESS
# ============================================================

def request_storage_access():
    """Open Android's All files access screen when required."""
    if not ANDROID_AVAILABLE:
        return

    try:
        Build_VERSION = autoclass("android.os.Build$VERSION")
        if Build_VERSION.SDK_INT < 30:
            return

        Environment = autoclass("android.os.Environment")
        if Environment.isExternalStorageManager():
            return

        PythonActivity = autoclass("org.kivy.android.PythonActivity")
        Intent = autoclass("android.content.Intent")
        Settings = autoclass("android.provider.Settings")
        Uri = autoclass("android.net.Uri")

        activity = PythonActivity.mActivity
        package_name = activity.getPackageName()

        intent = Intent(
            Settings.ACTION_MANAGE_APP_ALL_FILES_ACCESS_PERMISSION
        )
        intent.setData(Uri.parse("package:" + package_name))
        activity.startActivity(intent)

    except Exception as error:
        print("Storage permission screen could not be opened:", error)


# ============================================================
# SETTINGS
# ============================================================

WHATSAPP_FOLDER = (
    "/storage/emulated/0/Download/"
    "WhatsApp Chat with Hyperscale"
)

REPORT_FOLDER = (
    "/storage/emulated/0/Download/"
    "WhatsApp_Reports"
)

AUTO_SELECT_LATEST_ZIP = True

ZIP_NAME = "WhatsApp Chat with Hyperscale.zip"

CHAT_PREFIX = "WhatsApp Chat with"

MEDIA_FOLDER_NAME = "Media"
PREVIEW_FOLDER_NAME = "_Excel_Previews"


# ============================================================
# MEDIA TYPES
# ============================================================

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".bmp",
    ".webp"
}

VIDEO_EXTENSIONS = {
    ".mp4",
    ".3gp",
    ".mov",
    ".avi",
    ".mkv"
}

DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".ppt",
    ".pptx",
    ".zip",
    ".rar",
    ".txt"
}

AUDIO_EXTENSIONS = {
    ".mp3",
    ".m4a",
    ".aac",
    ".ogg",
    ".opus",
    ".wav"
}


# ============================================================
# PROCESS NAMES
# ============================================================

PROCESS_NAMES = [
    "Find WhatsApp ZIP",
    "Extract WhatsApp ZIP",
    "Find chat TXT",
    "Read chat messages",
    "Build media index",
    "Process messages",
    "Copy media files",
    "Create Excel previews",
    "Save Excel report",
    "Create Report ZIP",
    "Clean temporary files",
    "Completed",
]


# ============================================================
# ANDROID NOTIFICATION
# ============================================================

def request_notification_permission():

    if not ANDROID_AVAILABLE:
        return

    try:

        PythonActivity = autoclass(
            "org.kivy.android.PythonActivity"
        )

        Build_VERSION = autoclass(
            "android.os.Build$VERSION"
        )

        activity = PythonActivity.mActivity

        # Android 13+
        if Build_VERSION.SDK_INT >= 33:

            permission = (
                "android.permission.POST_NOTIFICATIONS"
            )

            try:

                activity.requestPermissions(
                    [permission],
                    1001
                )

                print(
                    "Notification permission requested."
                )

            except Exception as error:

                print(
                    "Could not request notification permission:",
                    error
                )

    except Exception as error:

        print(
            "Notification permission error:",
            error
        )


def send_export_notification():

    if not ANDROID_AVAILABLE:

        print(
            "Android notification unavailable."
        )

        return

    try:

        # ----------------------------------------------------
        # Android classes
        # ----------------------------------------------------

        PythonActivity = autoclass(
            "org.kivy.android.PythonActivity"
        )

        Context = autoclass(
            "android.content.Context"
        )

        Build_VERSION = autoclass(
            "android.os.Build$VERSION"
        )

        NotificationManager = autoclass(
            "android.app.NotificationManager"
        )

        NotificationChannel = autoclass(
            "android.app.NotificationChannel"
        )

        NotificationBuilder = autoclass(
            "android.app.Notification$Builder"
        )

        activity = PythonActivity.mActivity

        channel_id = (
            "whatsapp_export_channel"
        )

        # ----------------------------------------------------
        # Notification manager
        # ----------------------------------------------------

        notification_manager = (
            activity.getSystemService(
                Context.NOTIFICATION_SERVICE
            )
        )

        # ----------------------------------------------------
        # Android 8+ notification channel
        # ----------------------------------------------------

        if Build_VERSION.SDK_INT >= 26:

            channel = NotificationChannel(
                channel_id,
                "WhatsApp Chat Exporter",
                NotificationManager.IMPORTANCE_HIGH
            )

            channel.setDescription(
                "WhatsApp export completion notifications"
            )

            notification_manager.createNotificationChannel(
                channel
            )

        # ----------------------------------------------------
        # Notification builder
        # ----------------------------------------------------

        if Build_VERSION.SDK_INT >= 26:

            builder = NotificationBuilder(
                activity,
                channel_id
            )

        else:

            builder = NotificationBuilder(
                activity
            )

        # ----------------------------------------------------
        # Notification content
        # ----------------------------------------------------

        builder.setContentTitle(
            "WhatsApp Export Completed"
        )

        builder.setContentText(
            "Your WhatsApp chat report is ready."
        )

        builder.setSmallIcon(
            activity.getApplicationInfo().icon
        )

        builder.setAutoCancel(
            True
        )

        # ----------------------------------------------------
        # Show notification
        # ----------------------------------------------------

        notification_manager.notify(
            1001,
            builder.build()
        )

        print(
            "Android notification sent successfully."
        )

    except Exception as error:

        print(
            "Notification failed:",
            error
        )


# ============================================================
# STATUS ICON
# ============================================================

class StatusIcon(Widget):

    def __init__(self, **kwargs):

        super().__init__(**kwargs)

        self.state = "waiting"
        self.angle = 0

        self.bind(
            pos=self.redraw,
            size=self.redraw
        )

        Clock.schedule_interval(
            self.animate,
            1 / 20
        )

    def set_state(self, state):

        self.state = state

        self.redraw()

    def animate(self, dt):

        if self.state == "active":

            self.angle = (
                self.angle + 15
            ) % 360

            self.redraw()

    def redraw(self, *args):

        self.canvas.clear()

        with self.canvas:

            # ------------------------------------------------
            # COMPLETE
            # ------------------------------------------------

            if self.state == "complete":

                Color(
                    0.20,
                    0.85,
                    0.40,
                    1
                )

                Line(
                    circle=(
                        self.center_x,
                        self.center_y,
                        dp(10)
                    ),
                    width=2.2
                )

                Line(
                    points=[
                        self.center_x - dp(5),
                        self.center_y,

                        self.center_x - dp(1),
                        self.center_y - dp(5),

                        self.center_x + dp(7),
                        self.center_y + dp(6)
                    ],
                    width=2.6,
                    cap="round"
                )

            # ------------------------------------------------
            # ACTIVE
            # ------------------------------------------------

            elif self.state == "active":

                Color(
                    0.20,
                    0.65,
                    1,
                    1
                )

                Line(
                    circle=(
                        self.center_x,
                        self.center_y,
                        dp(10),
                        self.angle,
                        self.angle + 270
                    ),
                    width=3
                )

            # ------------------------------------------------
            # WAITING
            # ------------------------------------------------

            else:

                Color(
                    0.55,
                    0.55,
                    0.55,
                    1
                )

                Line(
                    circle=(
                        self.center_x,
                        self.center_y,
                        dp(9)
                    ),
                    width=1.7
                )


# ============================================================
# PROCESS ROW
# ============================================================

class ProcessRow(BoxLayout):

    def __init__(
        self,
        number,
        name,
        **kwargs
    ):

        super().__init__(
            orientation="vertical",
            size_hint_y=None,
            height=dp(57),
            spacing=dp(2),
            **kwargs
        )

        self.number = number
        self.name = name
        self.percent = 0

        top = BoxLayout(
            orientation="horizontal",
            size_hint_y=None,
            height=dp(30),
            spacing=dp(7)
        )

        self.icon = StatusIcon(
            size_hint_x=None,
            width=dp(28)
        )

        self.label = Label(
            text=f"{number}. {name}",
            font_size="14sp",
            halign="left",
            valign="middle"
        )

        self.label.bind(
            size=lambda obj, value:
            setattr(
                obj,
                "text_size",
                value
            )
        )

        self.percent_label = Label(
            text="0%",
            font_size="13sp",
            size_hint_x=None,
            width=dp(45),
            halign="right",
            valign="middle"
        )

        self.percent_label.bind(
            size=lambda obj, value:
            setattr(
                obj,
                "text_size",
                value
            )
        )

        top.add_widget(
            self.icon
        )

        top.add_widget(
            self.label
        )

        top.add_widget(
            self.percent_label
        )

        self.bar = ProgressBar(
            max=100,
            value=0,
            size_hint_y=None,
            height=dp(5)
        )

        self.add_widget(
            top
        )

        self.add_widget(
            self.bar
        )

    def set_state(
        self,
        state,
        percent
    ):

        self.percent = max(
            0,
            min(
                100,
                int(percent)
            )
        )

        self.bar.value = (
            self.percent
        )

        self.percent_label.text = (
            f"{self.percent}%"
        )

        self.icon.set_state(
            state
        )

        if state == "active":

            self.label.text = (
                f"{self.number}. "
                f"{self.name} - Working..."
            )

        else:

            self.label.text = (
                f"{self.number}. "
                f"{self.name}"
            )


# ============================================================
# WHATSAPP EXPORTER
# ============================================================

class WhatsAppExporter:

    def __init__(
        self,
        progress_callback
    ):

        self.progress_callback = (
            progress_callback
        )

    def progress(self, text):

        print(text)

        if self.progress_callback:

            self.progress_callback(
                text
            )

    # ========================================================
    # DELETE OLD REPORTS
    # ========================================================

    def delete_old_reports(self):

        os.makedirs(
            REPORT_FOLDER,
            exist_ok=True
        )

        for item in os.listdir(
            REPORT_FOLDER
        ):

            path = os.path.join(
                REPORT_FOLDER,
                item
            )

            try:

                if os.path.isfile(path):

                    os.remove(path)

                elif os.path.isdir(path):

                    shutil.rmtree(path)

            except Exception as error:

                print(
                    "Could not delete:",
                    path,
                    error
                )

    # ========================================================
    # FIND ZIP
    # ========================================================

    def find_zip_file(self):

        if not os.path.isdir(
            WHATSAPP_FOLDER
        ):

            raise FileNotFoundError(
                "WhatsApp folder was not found:\n\n"
                + WHATSAPP_FOLDER
            )

        zip_files = []

        for filename in os.listdir(
            WHATSAPP_FOLDER
        ):

            path = os.path.join(
                WHATSAPP_FOLDER,
                filename
            )

            if not os.path.isfile(path):

                continue

            if not filename.lower().endswith(
                ".zip"
            ):

                continue

            lower = filename.lower()

            if (
                "_whatsapp_report"
                in lower
                or
                "_report"
                in lower
            ):

                continue

            zip_files.append(
                path
            )

        if not zip_files:

            raise FileNotFoundError(
                "No original WhatsApp ZIP was found."
            )

        if AUTO_SELECT_LATEST_ZIP:

            return max(
                zip_files,
                key=os.path.getmtime
            )

        selected = os.path.join(
            WHATSAPP_FOLDER,
            ZIP_NAME
        )

        if not os.path.isfile(
            selected
        ):

            raise FileNotFoundError(
                "Specified ZIP was not found:\n\n"
                + selected
            )

        return selected

    # ========================================================
    # EXTRACT ZIP
    # ========================================================

    def extract_zip(
        self,
        zip_path,
        extract_folder
    ):

        if os.path.exists(
            extract_folder
        ):

            shutil.rmtree(
                extract_folder
            )

        os.makedirs(
            extract_folder,
            exist_ok=True
        )

        with zipfile.ZipFile(
            zip_path,
            "r"
        ) as archive:

            archive.extractall(
                extract_folder
            )

    # ========================================================
    # FIND CHAT TXT
    # ========================================================

    def find_chat_file(
        self,
        root_folder
    ):

        for root, _, files in os.walk(
            root_folder
        ):

            for filename in files:

                if (
                    filename.lower().endswith(
                        ".txt"
                    )
                    and
                    filename.lower().startswith(
                        CHAT_PREFIX.lower()
                    )
                ):

                    return os.path.join(
                        root,
                        filename
                    )

        return None

    # ========================================================
    # READ TXT
    # ========================================================

    def read_whatsapp_file(
        self,
        file_path
    ):

        encodings = (
            "utf-8-sig",
            "utf-8",
            "utf-16",
            "utf-16-le",
            "utf-16-be"
        )

        for encoding in encodings:

            try:

                with open(
                    file_path,
                    "r",
                    encoding=encoding
                ) as file:

                    return (
                        file.read()
                        .splitlines()
                    )

            except UnicodeError:

                continue

        raise ValueError(
            "Could not decode WhatsApp TXT file."
        )

    # ========================================================
    # MESSAGE PATTERN
    # ========================================================

    MESSAGE_PATTERN = re.compile(
        r"^\s*"
        r"\d{1,2}/\d{1,2}/\d{4},"
        r"\s*\d{1,2}:\d{2}"
        r"\s*-\s*"
    )

    # ========================================================
    # CHECK MESSAGE
    # ========================================================

    def is_whatsapp_message(
        self,
        line
    ):

        return bool(
            self.MESSAGE_PATTERN.match(
                line
            )
        )

    # ========================================================
    # PARSE MESSAGE
    # ========================================================

    def parse_whatsapp_message(
        self,
        line
    ):

        match = re.match(
            r"^\s*"
            r"(\d{1,2}/\d{1,2}/\d{4}),"
            r"\s*(\d{1,2}:\d{2})"
            r"\s*-\s*(.*)$",
            line
        )

        if not match:

            return (
                "",
                "",
                "",
                ""
            )

        date = (
            match.group(1)
            .strip()
        )

        time = (
            match.group(2)
            .strip()
        )

        body = (
            match.group(3)
            .strip()
        )

        colon = body.find(
            ": "
        )

        if colon >= 0:

            sender = (
                body[:colon]
                .strip()
            )

            message = (
                body[colon + 2:]
                .strip()
            )

        else:

            sender = ""

            message = body

        return (
            date,
            time,
            sender,
            message
        )

    # ========================================================
    # BUILD MEDIA INDEX
    # ========================================================

    def build_media_index(
        self,
        root_folder
    ):

        media_index = {}

        for root, _, files in os.walk(
            root_folder
        ):

            for filename in files:

                path = os.path.join(
                    root,
                    filename
                )

                media_index.setdefault(
                    filename.lower(),
                    path
                )

        return media_index

    # ========================================================
    # MEDIA FILENAME PATTERN
    # ========================================================

    MEDIA_FILENAME_PATTERN = re.compile(
        r"[\w.\- ()]+?\.(?:"
        r"jpg|jpeg|png|gif|webp|bmp|"
        r"mp4|3gp|mov|avi|mkv|"
        r"pdf|doc|docx|xls|xlsx|"
        r"ppt|pptx|zip|rar|txt|"
        r"mp3|m4a|aac|ogg|opus|wav|vcf"
        r")",
        re.IGNORECASE
    )

    # ========================================================
    # FIND MEDIA
    # ========================================================

    def find_media(
        self,
        message,
        media_index
    ):

        matches = (
            self.MEDIA_FILENAME_PATTERN
            .findall(message)
        )

        for match in matches:

            filename = match.strip()

            path = media_index.get(
                filename.lower()
            )

            if path:

                return (
                    filename,
                    path
                )

        return (
            "",
            ""
        )

    # ========================================================
    # MEDIA TYPE
    # ========================================================

    def get_media_type(
        self,
        filename,
        message
    ):

        if not filename:

            if (
                "<Media omitted>"
                in message
            ):

                return "Media"

            return "Text"

        ext = os.path.splitext(
            filename
        )[1].lower()

        if ext in IMAGE_EXTENSIONS:

            return "Image"

        if ext in VIDEO_EXTENSIONS:

            return "Video"

        if ext in DOCUMENT_EXTENSIONS:

            return "Document"

        if ext in AUDIO_EXTENSIONS:

            return "Audio"

        if ext == ".vcf":

            return "Contact"

        return "File"

    # ========================================================
    # COPY MEDIA
    # ========================================================

    def copy_media(
        self,
        source,
        destination_folder
    ):

        try:

            os.makedirs(
                destination_folder,
                exist_ok=True
            )

            filename = os.path.basename(
                source
            )

            destination = os.path.join(
                destination_folder,
                filename
            )

            if os.path.exists(
                destination
            ):

                name, extension = (
                    os.path.splitext(
                        filename
                    )
                )

                counter = 2

                while os.path.exists(
                    destination
                ):

                    destination = os.path.join(
                        destination_folder,
                        f"{name}_{counter}{extension}"
                    )

                    counter += 1

            shutil.copy2(
                source,
                destination
            )

            return destination

        except Exception as error:

            print(
                "Media copy failed:",
                error
            )

            return ""

    # ========================================================
    # CREATE IMAGE PREVIEW
    # ========================================================

    def create_preview(
        self,
        image_path,
        preview_folder
    ):

        if not PIL_AVAILABLE:

            return None

        try:

            os.makedirs(
                preview_folder,
                exist_ok=True
            )

            filename = os.path.basename(
                image_path
            )

            name = os.path.splitext(
                filename
            )[0]

            preview_path = os.path.join(
                preview_folder,
                name + "_preview.jpg"
            )

            image = Image.open(
                image_path
            )

            try:

                image.seek(0)

            except Exception:

                pass

            if image.mode != "RGB":

                image = image.convert(
                    "RGB"
                )

            image.thumbnail(
                (300, 300),
                Image.Resampling.LANCZOS
            )

            image.save(
                preview_path,
                "JPEG",
                quality=85
            )

            return preview_path

        except Exception as error:

            print(
                "Preview failed:",
                error
            )

            return None

    # ========================================================
    # IMPORT CHAT
    # ========================================================

    def import_chat(
        self,
        zip_path
    ):

        zip_name = os.path.splitext(
            os.path.basename(
                zip_path
            )
        )[0]

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        extract_folder = os.path.join(
            REPORT_FOLDER,
            zip_name +
            "_Extracted_" +
            timestamp
        )

        media_folder = os.path.join(
            REPORT_FOLDER,
            MEDIA_FOLDER_NAME
        )

        preview_folder = os.path.join(
            REPORT_FOLDER,
            PREVIEW_FOLDER_NAME
        )

        # ----------------------------------------------------
        # EXTRACT
        # ----------------------------------------------------

        self.progress(
            "Extracting WhatsApp ZIP..."
        )

        self.extract_zip(
            zip_path,
            extract_folder
        )

        self.progress(
            "Extraction completed"
        )

        # ----------------------------------------------------
        # FIND TXT
        # ----------------------------------------------------

        self.progress(
            "Finding chat TXT..."
        )

        chat_file = self.find_chat_file(
            extract_folder
        )

        if not chat_file:

            raise FileNotFoundError(
                "WhatsApp TXT file was not found."
            )

        self.progress(
            "Chat TXT found"
        )

        # ----------------------------------------------------
        # READ CHAT
        # ----------------------------------------------------

        self.progress(
            "Reading chat..."
        )

        lines = self.read_whatsapp_file(
            chat_file
        )

        self.progress(
            f"Chat lines loaded: {len(lines)}"
        )

        # ----------------------------------------------------
        # MEDIA INDEX
        # ----------------------------------------------------

        self.progress(
            "Building media index..."
        )

        media_index = (
            self.build_media_index(
                extract_folder
            )
        )

        self.progress(
            f"Media files indexed: "
            f"{len(media_index)}"
        )

        # ----------------------------------------------------
        # WORKBOOK
        # ----------------------------------------------------

        workbook = Workbook()

        worksheet = (
            workbook.active
        )

        worksheet.title = (
            "WhatsApp Chat"
        )

        headers = [
            "Date",
            "Time",
            "Name",
            "Message",
            "Media File",
            "Media Type",
            "Open File",
            "Preview"
        ]

        worksheet.append(
            headers
        )

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                vertical="center"
            )

        # ----------------------------------------------------
        # TOTAL MESSAGE COUNT
        # ----------------------------------------------------

        total_messages = sum(
            1
            for line in lines
            if self.is_whatsapp_message(
                line
            )
        )

        if total_messages <= 0:

            total_messages = 1

        current_date = ""
        current_time = ""
        current_sender = ""
        current_message = ""

        message_count = 0
        media_count = 0
        preview_count = 0

        # ----------------------------------------------------
        # WRITE MESSAGE
        # ----------------------------------------------------

        def write_message():

            nonlocal current_date
            nonlocal current_time
            nonlocal current_sender
            nonlocal current_message
            nonlocal message_count
            nonlocal media_count
            nonlocal preview_count

            if not current_message:

                return

            media_file, source_media = (
                self.find_media(
                    current_message,
                    media_index
                )
            )

            media_type = (
                self.get_media_type(
                    media_file,
                    current_message
                )
            )

            final_media = ""

            # ------------------------------------------------
            # COPY MEDIA
            # ------------------------------------------------

            if source_media:

                final_media = (
                    self.copy_media(
                        source_media,
                        media_folder
                    )
                )

                if final_media:

                    media_count += 1

            # ------------------------------------------------
            # WRITE EXCEL ROW
            # ------------------------------------------------

            worksheet.append([
                current_date,
                current_time,
                current_sender.strip(),
                current_message,
                media_file,
                media_type,
                "",
                ""
            ])

            excel_row = (
                worksheet.max_row
            )

            # ------------------------------------------------
            # HYPERLINK
            # ------------------------------------------------

            if final_media:

                cell = worksheet.cell(
                    excel_row,
                    7
                )

                cell.value = (
                    "Open File"
                )

                try:

                    relative = os.path.relpath(
                        final_media,
                        REPORT_FOLDER
                    ).replace(
                        os.sep,
                        "/"
                    )

                    cell.hyperlink = (
                        relative
                    )

                    cell.style = (
                        "Hyperlink"
                    )

                except Exception as error:

                    print(
                        "Hyperlink failed:",
                        error
                    )

                # --------------------------------------------
                # IMAGE PREVIEW
                # --------------------------------------------

                extension = (
                    os.path.splitext(
                        final_media
                    )[1].lower()
                )

                if extension in IMAGE_EXTENSIONS:

                    preview_path = (
                        self.create_preview(
                            final_media,
                            preview_folder
                        )
                    )

                    if preview_path:

                        try:

                            image = XLImage(
                                preview_path
                            )

                            max_size = 100

                            if image.width > max_size:

                                ratio = (
                                    max_size /
                                    image.width
                                )

                                image.width = int(
                                    image.width *
                                    ratio
                                )

                                image.height = int(
                                    image.height *
                                    ratio
                                )

                            if image.height > max_size:

                                ratio = (
                                    max_size /
                                    image.height
                                )

                                image.width = int(
                                    image.width *
                                    ratio
                                )

                                image.height = int(
                                    image.height *
                                    ratio
                                )

                            worksheet.add_image(
                                image,
                                f"H{excel_row}"
                            )

                            worksheet.row_dimensions[
                                excel_row
                            ].height = 85

                            preview_count += 1

                        except Exception as error:

                            print(
                                "Preview insert failed:",
                                error
                            )

            # ------------------------------------------------
            # MESSAGE COUNTER
            # ------------------------------------------------

            message_count += 1

            percent = int(
                message_count /
                max(total_messages, 1)
                * 100
            )

            percent = max(
                1,
                min(100, percent)
            )

            self.progress(
                f"Processed "
                f"{message_count}/"
                f"{total_messages} messages..."
            )

        # ----------------------------------------------------
        # PROCESS CHAT
        # ----------------------------------------------------

        self.progress(
            "Processing messages..."
        )

        for raw_line in lines:

            line = raw_line.replace(
                "\ufeff",
                ""
            )

            if not line.strip():

                if current_message:

                    current_message += "\n"

                continue

            if self.is_whatsapp_message(
                line
            ):

                write_message()

                (
                    current_date,
                    current_time,
                    current_sender,
                    current_message
                ) = (
                    self.parse_whatsapp_message(
                        line
                    )
                )

            else:

                if current_message:

                    current_message += (
                        "\n" + line
                    )

        write_message()

        # ----------------------------------------------------
        # COLUMN WIDTHS
        # ----------------------------------------------------

        widths = {
            "A": 13,
            "B": 12,
            "C": 28,
            "D": 60,
            "E": 40,
            "F": 15,
            "G": 15,
            "H": 18
        }

        for column, width in widths.items():

            worksheet.column_dimensions[
                column
            ].width = width

        # ----------------------------------------------------
        # CELL FORMATTING
        # ----------------------------------------------------

        for row in worksheet.iter_rows():

            for index in (
                0,
                1,
                2,
                5,
                6,
                7
            ):

                row[index].alignment = (
                    Alignment(
                        vertical="top"
                    )
                )

            row[3].alignment = (
                Alignment(
                    wrap_text=True,
                    vertical="top"
                )
            )

            row[4].alignment = (
                Alignment(
                    wrap_text=True,
                    vertical="top"
                )
            )

        worksheet.freeze_panes = (
            "A2"
        )

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        # ----------------------------------------------------
        # SAVE EXCEL
        # ----------------------------------------------------

        output_file = os.path.join(
            REPORT_FOLDER,
            zip_name +
            "_WhatsApp_Chat.xlsx"
        )

        if os.path.exists(
            output_file
        ):

            os.remove(
                output_file
            )

        self.progress(
            "Saving Excel report..."
        )

        workbook.save(
            output_file
        )

        self.progress(
            "Excel report saved"
        )

        return (
            output_file,
            chat_file,
            message_count,
            extract_folder,
            media_folder,
            media_count,
            preview_count
        )

    # ========================================================
    # CREATE REPORT ZIP
    # ========================================================

    def create_report_zip(
        self,
        output_excel,
        media_folder
    ):

        self.progress(
            "Creating report ZIP..."
        )

        output_directory = (
            os.path.dirname(
                output_excel
            )
        )

        report_zip = os.path.join(
            output_directory,
            "Share_WhatsApp_Report.zip"
        )

        if os.path.exists(
            report_zip
        ):

            os.remove(
                report_zip
            )

        with zipfile.ZipFile(
            report_zip,
            "w",
            compression=zipfile.ZIP_DEFLATED
        ) as archive:

            archive.write(
                output_excel,
                arcname=os.path.basename(
                    output_excel
                )
            )

            if os.path.isdir(
                media_folder
            ):

                for root, _, files in os.walk(
                    media_folder
                ):

                    for filename in files:

                        full_path = os.path.join(
                            root,
                            filename
                        )

                        relative = os.path.relpath(
                            full_path,
                            output_directory
                        ).replace(
                            os.sep,
                            "/"
                        )

                        archive.write(
                            full_path,
                            arcname=relative
                        )

        self.progress(
            "Report ZIP created"
        )

        return report_zip

    # ========================================================
    # CLEANUP
    # ========================================================

    def cleanup(
        self,
        extract_folder
    ):

        self.progress(
            "Cleaning temporary files..."
        )

        try:

            if os.path.exists(
                extract_folder
            ):

                shutil.rmtree(
                    extract_folder
                )

        except Exception as error:

            print(
                "Cleanup failed:",
                error
            )

        self.progress(
            "Temporary files cleaned"
        )

    # ========================================================
    # MAIN EXPORT
    # ========================================================

    def run(self):

        # ----------------------------------------------------
        # CLEAN OLD REPORTS
        # ----------------------------------------------------

        self.delete_old_reports()

        # ----------------------------------------------------
        # FIND ZIP
        # ----------------------------------------------------

        self.progress(
            "Searching for WhatsApp ZIP..."
        )

        zip_path = (
            self.find_zip_file()
        )

        self.progress(
            "Original WhatsApp ZIP found: "
            + os.path.basename(
                zip_path
            )
        )

        # ----------------------------------------------------
        # IMPORT CHAT
        # ----------------------------------------------------

        (
            output_file,
            chat_file,
            message_count,
            extract_folder,
            media_folder,
            media_count,
            preview_count
        ) = self.import_chat(
            zip_path
        )

        # ----------------------------------------------------
        # CREATE REPORT ZIP
        # ----------------------------------------------------

        report_zip = (
            self.create_report_zip(
                output_file,
                media_folder
            )
        )

        # ----------------------------------------------------
        # CLEANUP
        # ----------------------------------------------------

        self.cleanup(
            extract_folder
        )

        # ----------------------------------------------------
        # COMPLETE
        # ----------------------------------------------------

        self.progress(
            "Export completed"
        )

        return {
            "zip_path": zip_path,
            "chat_file": chat_file,
            "output_file": output_file,
            "media_folder": media_folder,
            "report_zip": report_zip,
            "messages": message_count,
            "media": media_count,
            "previews": preview_count
        }


# ============================================================
# KIVY APPLICATION
# ============================================================

class WhatsAppExporterApp(App):

    def build(self):

        self.title = (
            "WhatsApp Chat Exporter"
        )

        # Android 11+ access for ZIP/TXT/media/Excel files.
        Clock.schedule_once(
            lambda dt: request_storage_access(),
            2
        )

        # ----------------------------------------------------
        # REQUEST NOTIFICATION PERMISSION
        # ----------------------------------------------------

        Clock.schedule_once(
            lambda dt:
            request_notification_permission(),
            1
        )

        root = BoxLayout(
            orientation="vertical",
            padding=dp(12),
            spacing=dp(5)
        )

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = Label(
            text="WHATSAPP CHAT EXPORTER",
            font_size="22sp",
            bold=True,
            size_hint_y=None,
            height=dp(38)
        )

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        self.status = Label(
            text="Ready",
            font_size="15sp",
            size_hint_y=None,
            height=dp(28)
        )

        # ----------------------------------------------------
        # OVERALL PROGRESS
        # ----------------------------------------------------

        self.overall_bar = ProgressBar(
            max=100,
            value=0,
            size_hint_y=None,
            height=dp(8)
        )

        self.overall_label = Label(
            text="Overall progress: 0%",
            font_size="13sp",
            size_hint_y=None,
            height=dp(22)
        )

        # ----------------------------------------------------
        # PROCESS LIST
        # ----------------------------------------------------

        process_box = BoxLayout(
            orientation="vertical",
            size_hint_y=None,
            spacing=dp(2)
        )

        self.rows = []

        for number, name in enumerate(
            PROCESS_NAMES,
            1
        ):

            row = ProcessRow(
                number,
                name
            )

            self.rows.append(
                row
            )

            process_box.add_widget(
                row
            )

        process_box.height = sum(
            row.height
            for row in self.rows
        )

        scroll = ScrollView(
            do_scroll_x=False
        )

        scroll.add_widget(
            process_box
        )

        # ----------------------------------------------------
        # EXPORT BUTTON
        # ----------------------------------------------------

        self.export_button = Button(
            text="START EXPORT",
            font_size="18sp",
            size_hint_y=None,
            height=dp(56)
        )

        self.export_button.bind(
            on_press=self.start_export
        )

        # ----------------------------------------------------
        # ADD WIDGETS
        # ----------------------------------------------------

        root.add_widget(
            title
        )

        root.add_widget(
            self.status
        )

        root.add_widget(
            self.overall_bar
        )

        root.add_widget(
            self.overall_label
        )

        root.add_widget(
            scroll
        )

        root.add_widget(
            self.export_button
        )

        return root

    # ========================================================
    # SET PROCESS
    # ========================================================

    def set_process(
        self,
        number,
        state,
        percent
    ):

        self.rows[
            number - 1
        ].set_state(
            state,
            percent
        )

        self.update_overall()

    # ========================================================
    # UPDATE OVERALL
    # ========================================================

    def update_overall(self):

        total = sum(
            row.percent
            for row in self.rows
        )

        overall = int(
            total /
            len(self.rows)
        )

        self.overall_bar.value = (
            overall
        )

        self.overall_label.text = (
            f"Overall progress: "
            f"{overall}%"
        )

    # ========================================================
    # RESET
    # ========================================================

    def reset(self):

        for row in self.rows:

            row.set_state(
                "waiting",
                0
            )

        self.overall_bar.value = 0

        self.overall_label.text = (
            "Overall progress: 0%"
        )

    # ========================================================
    # START EXPORT
    # ========================================================

    def start_export(
        self,
        instance
    ):

        if self.export_button.disabled:

            return

        self.reset()

        self.export_button.disabled = (
            True
        )

        self.export_button.text = (
            "EXPORTING..."
        )

        self.status.text = (
            "Starting export..."
        )

        self.set_process(
            1,
            "active",
            0
        )

        Thread(
            target=self.export_worker,
            daemon=True
        ).start()

    # ========================================================
    # EXPORT WORKER
    # ========================================================

    def export_worker(self):

        try:

            exporter = WhatsAppExporter(
                self.set_status
            )

            result = exporter.run()

            Clock.schedule_once(
                lambda dt:
                self.export_success(
                    result
                )
            )

        except Exception as error:

            Clock.schedule_once(
                lambda dt:
                self.export_failed(
                    str(error)
                )
            )

    # ========================================================
    # STATUS CALLBACK
    # ========================================================

    def set_status(
        self,
        text
    ):

        Clock.schedule_once(
            lambda dt:
            self.handle_status(
                text
            )
        )

    # ========================================================
    # HANDLE STATUS
    # ========================================================

    def handle_status(
        self,
        text
    ):

        self.status.text = text

        t = text.lower()

        # ----------------------------------------------------
        # 1. FIND ZIP
        # ----------------------------------------------------

        if (
            "searching for whatsapp zip"
            in t
        ):

            self.set_process(
                1,
                "active",
                50
            )

        elif (
            "original whatsapp zip found"
            in t
        ):

            self.set_process(
                1,
                "complete",
                100
            )

            self.set_process(
                2,
                "active",
                0
            )

        # ----------------------------------------------------
        # 2. EXTRACT
        # ----------------------------------------------------

        elif (
            "extracting whatsapp zip"
            in t
        ):

            self.set_process(
                2,
                "active",
                50
            )

        elif (
            "extraction completed"
            in t
        ):

            self.set_process(
                2,
                "complete",
                100
            )

            self.set_process(
                3,
                "active",
                0
            )

        # ----------------------------------------------------
        # 3. CHAT TXT
        # ----------------------------------------------------

        elif (
            "finding chat txt"
            in t
        ):

            self.set_process(
                3,
                "active",
                50
            )

        elif (
            "chat txt found"
            in t
        ):

            self.set_process(
                3,
                "complete",
                100
            )

            self.set_process(
                4,
                "active",
                0
            )

        # ----------------------------------------------------
        # 4. READ CHAT
        # ----------------------------------------------------

        elif (
            "reading chat"
            in t
        ):

            self.set_process(
                4,
                "active",
                50
            )

        elif (
            "chat lines loaded"
            in t
        ):

            self.set_process(
                4,
                "complete",
                100
            )

            self.set_process(
                5,
                "active",
                0
            )

        # ----------------------------------------------------
        # 5. MEDIA INDEX
        # ----------------------------------------------------

        elif (
            "building media index"
            in t
        ):

            self.set_process(
                5,
                "active",
                50
            )

        elif (
            "media files indexed"
            in t
        ):

            self.set_process(
                5,
                "complete",
                100
            )

            self.set_process(
                6,
                "active",
                0
            )

        # ----------------------------------------------------
        # 6. MESSAGE PROCESSING
        # ----------------------------------------------------

        elif (
            "processed"
            in t
            and
            "messages"
            in t
        ):

            match = re.search(
                r"processed\s+(\d+)\s*/\s*(\d+)\s+messages",
                t
            )

            if match:

                current = int(
                    match.group(1)
                )

                total = int(
                    match.group(2)
                )

                percent = int(
                    current /
                    max(total, 1)
                    * 100
                )

                percent = max(
                    1,
                    min(100, percent)
                )

                # Process messages
                self.set_process(
                    6,
                    "active",
                    percent
                )

                # Copy media
                self.set_process(
                    7,
                    "active",
                    percent
                )

                # Create previews
                self.set_process(
                    8,
                    "active",
                    percent
                )

        # ----------------------------------------------------
        # 9. SAVE EXCEL
        # ----------------------------------------------------

        elif (
            "saving excel report"
            in t
        ):

            self.set_process(
                6,
                "complete",
                100
            )

            self.set_process(
                7,
                "complete",
                100
            )

            self.set_process(
                8,
                "complete",
                100
            )

            self.set_process(
                9,
                "active",
                50
            )

        elif (
            "excel report saved"
            in t
        ):

            self.set_process(
                9,
                "complete",
                100
            )

            self.set_process(
                10,
                "active",
                0
            )

        # ----------------------------------------------------
        # 10. REPORT ZIP
        # ----------------------------------------------------

        elif (
            "creating report zip"
            in t
        ):

            self.set_process(
                10,
                "active",
                50
            )

        elif (
            "report zip created"
            in t
        ):

            self.set_process(
                10,
                "complete",
                100
            )

            self.set_process(
                11,
                "active",
                0
            )

        # ----------------------------------------------------
        # 11. CLEANUP
        # ----------------------------------------------------

        elif (
            "cleaning temporary files"
            in t
        ):

            self.set_process(
                11,
                "active",
                50
            )

        elif (
            "temporary files cleaned"
            in t
        ):

            self.set_process(
                11,
                "complete",
                100
            )

            self.set_process(
                12,
                "active",
                50
            )

        # ----------------------------------------------------
        # 12. COMPLETE
        # ----------------------------------------------------

        elif (
            "export completed"
            in t
        ):

            self.set_process(
                12,
                "complete",
                100
            )

    # ========================================================
    # EXPORT SUCCESS
    # ========================================================

    def export_success(
        self,
        result
    ):

        # ----------------------------------------------------
        # MAKE EVERYTHING 100%
        # ----------------------------------------------------

        for row in self.rows:

            row.set_state(
                "complete",
                100
            )

        self.overall_bar.value = 100

        self.overall_label.text = (
            "Overall progress: 100%"
        )

        self.status.text = (
            "EXPORT COMPLETED"
        )

        # ----------------------------------------------------
        # SEND ANDROID NOTIFICATION
        # ----------------------------------------------------

        send_export_notification()

        # ----------------------------------------------------
        # ENABLE BUTTON
        # ----------------------------------------------------

        self.export_button.disabled = (
            False
        )

        self.export_button.text = (
            "START EXPORT"
        )

        # ----------------------------------------------------
        # CONSOLE INFORMATION
        # ----------------------------------------------------

        print(
            "================================"
        )

        print(
            "WHATSAPP EXPORT COMPLETED"
        )

        print(
            "================================"
        )

        print(
            "Excel:",
            result["output_file"]
        )

        print(
            "ZIP:",
            result["report_zip"]
        )

        print(
            "Messages:",
            result["messages"]
        )

        print(
            "Media:",
            result["media"]
        )

        print(
            "Previews:",
            result["previews"]
        )

        print(
            "================================"
        )

    # ========================================================
    # EXPORT FAILED
    # ========================================================

    def export_failed(
        self,
        error
    ):

        self.status.text = (
            "EXPORT FAILED"
        )

        self.export_button.disabled = (
            False
        )

        self.export_button.text = (
            "TRY AGAIN"
        )

        print(
            "================================"
        )

        print(
            "EXPORT ERROR:"
        )

        print(
            error
        )

        print(
            "================================"
        )


# ============================================================
# START APPLICATION
# ============================================================

if __name__ == "__main__":

    WhatsAppExporterApp().run()
