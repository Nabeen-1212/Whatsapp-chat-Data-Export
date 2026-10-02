import os
import re
import zipfile
import shutil
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


# ============================================================
# ANDROID / PYDROID SETTINGS
# ============================================================

WHATSAPP_FOLDER = "/storage/emulated/0/Download/WhatsApp Chat with Hyperscale"
REPORT_FOLDER = "/storage/emulated/0/Download/WhatsApp_Reports"

AUTO_SELECT_LATEST_ZIP = True
ZIP_NAME = "WhatsApp Chat with Hyperscale.zip"

CHAT_PREFIX = "WhatsApp Chat with"

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".3gp", ".mov", ".avi", ".mkv"}
DOCUMENT_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".zip", ".rar", ".txt"
}
AUDIO_EXTENSIONS = {".mp3", ".m4a", ".aac", ".ogg", ".opus", ".wav"}

MEDIA_FOLDER_NAME = "Media"
PREVIEW_FOLDER_NAME = "_Excel_Previews"


def delete_old_reports():
    os.makedirs(REPORT_FOLDER, exist_ok=True)

    for item in os.listdir(REPORT_FOLDER):
        full_path = os.path.join(REPORT_FOLDER, item)
        try:
            if os.path.isfile(full_path):
                os.remove(full_path)
            elif os.path.isdir(full_path):
                shutil.rmtree(full_path)
        except Exception as error:
            print("Could not delete:", full_path, error)


def find_zip_file(folder):
    if not os.path.isdir(folder):
        raise FileNotFoundError(
            "WhatsApp folder was not found:\n\n" + folder
        )

    zip_files = []

    for filename in os.listdir(folder):
        full_path = os.path.join(folder, filename)

        if not os.path.isfile(full_path):
            continue
        if not filename.lower().endswith(".zip"):
            continue

        lower_name = filename.lower()

        if "_whatsapp_report" in lower_name or "_report" in lower_name:
            continue

        zip_files.append(full_path)

    if not zip_files:
        raise FileNotFoundError(
            "No ORIGINAL WhatsApp ZIP was found in:\n\n" + folder
        )

    if AUTO_SELECT_LATEST_ZIP:
        return max(zip_files, key=os.path.getmtime)

    selected = os.path.join(folder, ZIP_NAME)

    if not os.path.isfile(selected):
        raise FileNotFoundError(
            "Specified ZIP file was not found:\n\n" + selected
        )

    return selected


def find_chat_file(root_folder):
    for root, _, files in os.walk(root_folder):
        for filename in files:
            if (
                filename.lower().endswith(".txt")
                and filename.lower().startswith(CHAT_PREFIX.lower())
            ):
                return os.path.join(root, filename)
    return None


def read_whatsapp_file(file_path):
    for encoding in (
        "utf-8-sig", "utf-8", "utf-16",
        "utf-16-le", "utf-16-be"
    ):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read().splitlines()
        except UnicodeError:
            continue

    raise ValueError("Could not decode the WhatsApp TXT file.")


MESSAGE_PATTERN = re.compile(
    r"^\s*\d{1,2}/\d{1,2}/\d{4},\s*\d{1,2}:\d{2}\s*-\s*"
)


def is_whatsapp_message(line):
    return bool(MESSAGE_PATTERN.match(line))


def parse_whatsapp_message(line):
    match = re.match(
        r"^\s*(\d{1,2}/\d{1,2}/\d{4}),\s*(\d{1,2}:\d{2})\s*-\s*(.*)$",
        line
    )

    if not match:
        return "", "", "", ""

    msg_date = match.group(1).strip()
    msg_time = match.group(2).strip()
    body = match.group(3).strip()

    colon_pos = body.find(": ")

    if colon_pos >= 0:
        sender = body[:colon_pos].strip()
        message = body[colon_pos + 2:].strip()
    else:
        sender = ""
        message = body

    return msg_date, msg_time, sender, message


def get_user_name(sender):
    return sender.strip()


def build_media_index(root_folder):
    media_index = {}

    for root, _, files in os.walk(root_folder):
        for filename in files:
            full_path = os.path.join(root, filename)
            media_index.setdefault(filename.lower(), full_path)

    return media_index


MEDIA_FILENAME_PATTERN = re.compile(
    r"[\w.\- ()]+?\.(?:"
    r"jpg|jpeg|png|gif|webp|bmp|"
    r"mp4|3gp|mov|avi|mkv|"
    r"pdf|doc|docx|xls|xlsx|ppt|pptx|"
    r"zip|rar|txt|"
    r"mp3|m4a|aac|ogg|opus|wav|vcf"
    r")",
    re.IGNORECASE
)


def find_media_from_message(message, media_index):
    matches = MEDIA_FILENAME_PATTERN.findall(message)

    for match in matches:
        filename = match.strip()
        path = media_index.get(filename.lower())

        if path:
            return filename, path

    return "", ""


def get_media_type(filename, message):
    if not filename:
        return "Media" if "<Media omitted>" in message else "Text"

    ext = os.path.splitext(filename)[1].lower()

    if ext in IMAGE_EXTENSIONS:
        return "Image"
    if ext in VIDEO_EXTENSIONS:
        return "Video"
    if ext in DOCUMENT_EXTENSIONS:
        return "Document"
    if ext in AUDIO_EXTENSIONS:
        return "Audio"
    if ext == ".vcf":
        return "Contact"

    return "File"


def create_preview_image(image_path, preview_folder):
    if not PIL_AVAILABLE:
        return None

    try:
        os.makedirs(preview_folder, exist_ok=True)

        original_name = os.path.basename(image_path)
        name = os.path.splitext(original_name)[0]
        preview_path = os.path.join(
            preview_folder, name + "_preview.jpg"
        )

        image = Image.open(image_path)

        try:
            image.seek(0)
        except Exception:
            pass

        if image.mode != "RGB":
            image = image.convert("RGB")

        image.thumbnail((300, 300), Image.Resampling.LANCZOS)
        image.save(preview_path, "JPEG", quality=85)

        return preview_path

    except Exception as error:
        print("Preview creation failed:", image_path, error)
        return None


def copy_media_to_report(media_path, media_folder):
    try:
        os.makedirs(media_folder, exist_ok=True)

        filename = os.path.basename(media_path)
        destination = os.path.join(media_folder, filename)

        if os.path.exists(destination):
            name, extension = os.path.splitext(filename)
            counter = 2

            while os.path.exists(destination):
                destination = os.path.join(
                    media_folder,
                    f"{name}_{counter}{extension}"
                )
                counter += 1

        shutil.copy2(media_path, destination)
        return destination

    except Exception as error:
        print("Could not copy media:", media_path, error)
        return ""


def extract_zip(zip_path, extract_folder):
    if os.path.exists(extract_folder):
        shutil.rmtree(extract_folder)

    os.makedirs(extract_folder, exist_ok=True)

    with zipfile.ZipFile(zip_path, "r") as z:
        z.extractall(extract_folder)


def import_whatsapp_zip(zip_path, progress_callback=None):
    def progress(text):
        print(text)
        if progress_callback:
            progress_callback(text)

    zip_name = os.path.splitext(os.path.basename(zip_path))[0]

    os.makedirs(REPORT_FOLDER, exist_ok=True)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    extract_folder = os.path.join(
        REPORT_FOLDER,
        zip_name + "_Extracted_" + stamp
    )

    media_folder = os.path.join(
        REPORT_FOLDER,
        MEDIA_FOLDER_NAME
    )

    preview_folder = os.path.join(
        REPORT_FOLDER,
        PREVIEW_FOLDER_NAME
    )

    os.makedirs(media_folder, exist_ok=True)
    os.makedirs(preview_folder, exist_ok=True)

    progress("Extracting WhatsApp ZIP...")
    extract_zip(zip_path, extract_folder)

    progress("Finding chat TXT...")
    chat_file = find_chat_file(extract_folder)

    if not chat_file:
        raise FileNotFoundError(
            "Could not find the WhatsApp TXT file.\n\n"
            "Expected filename beginning with:\n"
            "WhatsApp Chat with"
        )

    progress("Reading chat...")
    lines = read_whatsapp_file(chat_file)

    progress(f"Chat lines loaded: {len(lines)}")
    media_index = build_media_index(extract_folder)
    progress(f"Media files indexed: {len(media_index)}")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WhatsApp Chat"

    headers = [
        "Date", "Time", "Name", "Message",
        "Media File", "Media Type", "Open File", "Preview"
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    current_date = ""
    current_time = ""
    current_sender = ""
    current_message = ""

    message_count = 0
    media_count = 0
    preview_count = 0

    def write_message():
        nonlocal current_date, current_time, current_sender
        nonlocal current_message, message_count
        nonlocal media_count, preview_count

        if not current_message:
            return

        user_name = get_user_name(current_sender)

        media_file, original_media_path = find_media_from_message(
            current_message, media_index
        )

        media_type = get_media_type(
            media_file, current_message
        )

        final_media_path = ""

        if original_media_path:
            final_media_path = copy_media_to_report(
                original_media_path, media_folder
            )

            if final_media_path:
                media_count += 1

        worksheet.append([
            current_date,
            current_time,
            user_name,
            current_message,
            media_file,
            media_type,
            "",
            ""
        ])

        excel_row = worksheet.max_row

        if final_media_path:
            link_cell = worksheet.cell(excel_row, 7)
            link_cell.value = "Open File"

            try:
                relative_path = os.path.relpath(
                    final_media_path,
                    REPORT_FOLDER
                ).replace(os.sep, "/")

                link_cell.hyperlink = relative_path
                link_cell.style = "Hyperlink"
            except Exception as error:
                print("Could not create hyperlink:", error)

            extension = os.path.splitext(
                final_media_path
            )[1].lower()

            if extension in IMAGE_EXTENSIONS:
                try:
                    preview_path = create_preview_image(
                        final_media_path,
                        preview_folder
                    )

                    if preview_path:
                        image = XLImage(preview_path)

                        if image.width > 100:
                            ratio = 100 / image.width
                            image.width = int(image.width * ratio)
                            image.height = int(image.height * ratio)

                        if image.height > 100:
                            ratio = 100 / image.height
                            image.width = int(image.width * ratio)
                            image.height = int(image.height * ratio)

                        worksheet.add_image(
                            image, f"H{excel_row}"
                        )

                        worksheet.row_dimensions[
                            excel_row
                        ].height = 85

                        preview_count += 1

                except Exception as error:
                    print(
                        "Excel preview failed:",
                        final_media_path,
                        error
                    )

        message_count += 1

        if message_count % 100 == 0:
            progress(f"Processed {message_count} messages...")

    for raw_line in lines:
        line = raw_line.replace("\ufeff", "")

        if not line.strip():
            if current_message:
                current_message += "\n"
            continue

        if is_whatsapp_message(line):
            write_message()

            (
                current_date,
                current_time,
                current_sender,
                current_message
            ) = parse_whatsapp_message(line)
        else:
            if current_message:
                current_message += "\n" + line

    write_message()

    column_widths = {
        "A": 13, "B": 12, "C": 28, "D": 60,
        "E": 40, "F": 15, "G": 15, "H": 18
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    for row in worksheet.iter_rows():
        for index in (0, 1, 2, 5, 6, 7):
            row[index].alignment = Alignment(vertical="top")

        row[3].alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        row[4].alignment = Alignment(
            wrap_text=True, vertical="top"
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    output_file = os.path.join(
        REPORT_FOLDER,
        zip_name + "_WhatsApp_Chat.xlsx"
    )

    if os.path.exists(output_file):
        os.remove(output_file)

    progress("Saving Excel report...")
    workbook.save(output_file)

    return (
        output_file,
        chat_file,
        message_count,
        extract_folder,
        media_folder,
        media_count,
        preview_count
    )


def create_report_zip(output_excel, media_folder):
    output_directory = os.path.dirname(output_excel)

    report_zip = os.path.join(
        output_directory,
        "Share_WhatsApp_Report.zip"
    )

    if os.path.exists(report_zip):
        os.remove(report_zip)

    with zipfile.ZipFile(
        report_zip,
        "w",
        compression=zipfile.ZIP_DEFLATED
    ) as archive:

        archive.write(
            output_excel,
            arcname=os.path.basename(output_excel)
        )

        if os.path.isdir(media_folder):
            for root, _, files in os.walk(media_folder):
                for filename in files:
                    full_path = os.path.join(root, filename)

                    relative_path = os.path.relpath(
                        full_path,
                        output_directory
                    ).replace(os.sep, "/")

                    archive.write(
                        full_path,
                        arcname=relative_path
                    )

    return report_zip


def delete_extracted_folder(extract_folder):
    try:
        if os.path.exists(extract_folder):
            shutil.rmtree(extract_folder)
    except Exception as error:
        print("Could not delete temporary folder:", error)


def run_export(progress_callback=None):
    """
    Main function called by the Kivy application.
    Returns a dictionary containing the result.
    """

    delete_old_reports()

    if progress_callback:
        progress_callback("Searching for WhatsApp ZIP...")

    zip_path = find_zip_file(WHATSAPP_FOLDER)

    if progress_callback:
        progress_callback(
            "ZIP found: " + os.path.basename(zip_path)
        )

    (
        output_file,
        chat_file,
        count,
        extract_folder,
        media_folder,
        media_count,
        preview_count
    ) = import_whatsapp_zip(
        zip_path,
        progress_callback
    )

    if progress_callback:
        progress_callback("Creating report ZIP...")

    report_zip = create_report_zip(
        output_file,
        media_folder
    )

    delete_extracted_folder(extract_folder)

    return {
        "zip_path": zip_path,
        "chat_file": chat_file,
        "output_file": output_file,
        "media_folder": media_folder,
        "report_zip": report_zip,
        "messages": count,
        "media": media_count,
        "previews": preview_count,
    }


if __name__ == "__main__":
    result = run_export(print)
    print("\nSUCCESS")
    print("Messages:", result["messages"])
    print("Media:", result["media"])
    print("Previews:", result["previews"])
    print("Excel:", result["output_file"])
    print("ZIP:", result["report_zip"])
